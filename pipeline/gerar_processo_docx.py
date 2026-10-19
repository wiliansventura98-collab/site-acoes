"""Gera a versão em Word (PROCESSO.docx) do documento de processo (PROCESSO.md).

Uso:
    python pipeline/gerar_processo_docx.py

Requer pandoc e python-docx. Rode sempre que PROCESSO.md for atualizado, para que as duas versões fiquem iguais.
"""
from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
TINTA, AZUL, LARANJA, LINHA, FUNDO = RGBColor(0x13, 0x29, 0x4B), RGBColor(0x25, 0x69, 0xAD), RGBColor(0xC4, 0x5A, 0x00), "C9D6E3", "EAF3FB"
FONTE = "Calibri"


def modelo(caminho: Path) -> None:
    """Documento de referência do pandoc com a tipografia e as cores do DATALUTA."""
    with open(caminho, "wb") as fh:
        fh.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout)
    d = Document(caminho)
    for s in d.styles:
        try:
            f = s.font
        except AttributeError:
            continue
        f.name = FONTE
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONTE)
    sty = lambda nome: next((x for x in d.styles if x.name == nome), None)  # pandoc grava os nomes como aparecem no Word
    sty("Normal").font.size = Pt(10.5)
    for nome, tam, cor in (("Title", 22, TINTA), ("Heading 1", 16, TINTA), ("Heading 2", 13, AZUL), ("Heading 3", 11.5, LARANJA)):
        st = sty(nome); st.font.size = Pt(tam); st.font.bold = True; st.font.color.rgb = cor
        st.paragraph_format.space_before = Pt(14 if nome == "Heading 1" else 10); st.paragraph_format.space_after = Pt(4)
    for nome in ("Body Text", "First Paragraph", "Compact"):
        if sty(nome): sty(nome).paragraph_format.space_after = Pt(5)
    for nome in ("Source Code", "Verbatim Char"):
        if sty(nome): sty(nome).font.name = "Consolas"; sty(nome).font.size = Pt(8.5)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2.2))
    for m in ("top_margin", "bottom_margin"): setattr(sec, m, Cm(2))
    d.save(caminho)


def bordas(tab) -> None:
    tbl = tab._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), LINHA); b.append(e)
    tblPr.append(b)


def acabamento(caminho: Path) -> None:
    d = Document(caminho)
    for tab in d.tables:
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordas(tab)
        w = tab._tbl.tblPr.find(qn("w:tblW"))  # tabela na largura da página
        if w is None:
            w = OxmlElement("w:tblW"); tab._tbl.tblPr.append(w)
        w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")
        for i, linha in enumerate(tab.rows):
            for cel in linha.cells:
                for p in cel.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == 0: r.font.bold = True
                if i == 0:
                    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), FUNDO)
                    cel._tc.get_or_add_tcPr().append(sh)
    # rodapé com o nome do documento
    rod = d.sections[0].footer.paragraphs[0]
    rod.text = "DATALUTA Argentina 2024 · Processo de construção da base"
    for r in rod.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x7C, 0x8A, 0xA0)
    d.save(caminho)


def main() -> None:
    md, saida = RAIZ / "PROCESSO.md", RAIZ / "PROCESSO.docx"
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "referencia.docx"
        modelo(ref)
        texto = md.read_text(encoding="utf-8")
        entrada = Path(tmp) / "processo.md"
        entrada.write_text(texto, encoding="utf-8")  # o sumário do próprio documento tem links internos que funcionam no Word
        subprocess.run(["pandoc", str(entrada), "-f", "gfm", "-o", str(saida), "--reference-doc", str(ref),
                        "-M", "lang=pt-BR"], check=True)
    acabamento(saida)
    print("Gerado:", saida)


if __name__ == "__main__":
    main()
